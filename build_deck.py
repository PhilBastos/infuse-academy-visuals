"""Build the Google-Slides-ready deck: clean PNG backgrounds + live Asap text boxes.

Coordinates are the 1920x1080 CSS pixels used in index.html (1px = 6350 EMU, 1px = 0.5pt).
Run after exporting png/0N-clean.png and png/0N-full.png.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

PX = 6350
NAVY, ORANGE, MINT, WHITE = '021A36', 'FF4F28', '78FFD6', 'FFFFFF'
GREY = '8993A0'

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)


def run(text, size, color, bold=False, ls=0.0):
    return dict(text=text, size=size, color=color, bold=bold, ls=ls)


def text(slide, x, y, w, h, lines, lh=1.15, align=PP_ALIGN.LEFT, wrap=False):
    """lines: list of paragraphs, each a list of runs."""
    tb = slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(max(r['size'] for r in runs) * 0.5 * lh)
        for r in runs:
            rr = p.add_run()
            rr.text = r['text']
            f = rr.font
            f.name, f.size, f.bold = 'Asap', Pt(r['size'] * 0.5), r['bold']
            f.color.rgb = RGBColor.from_string(r['color'])
            if r['ls']:
                rr._r.get_or_add_rPr().set('spc', str(round(r['ls'] * r['size'] * 0.5 * 100)))
    return tb


def dash(slide, x, y, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x * PX), Emu(y * PX), Emu(40 * PX), Emu(4 * PX))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()


def new_slide(bg, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(bg, 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = notes
    return s


TOOLS = ('Tools: character generated in Magnific (slide 1); visuals built in code (HTML, SVG, Canvas) with Claude (Anthropic); '
         'typography Asap; INFUSE Academy palette. Animated version: see the live link.')

# ---------- 01 ----------
s = new_slide('png/01-clean.png', TOOLS)
dash(s, 130, 100, ORANGE)
text(s, 186, 90, 500, 30, [[run('01  FIRST STEP', 22, ORANGE, True, .22)]])
text(s, 126, 150, 900, 430, [
    [run('No matter how', 104, NAVY, True, -.028)],
    [run('hard it looks,', 104, NAVY, True, -.028)],
    [run('there is always', 104, NAVY, True, -.028)],
    [run('a beginning.', 104, ORANGE, True, -.028)],
], lh=1.02)
text(s, 130, 1018, 500, 24, [[run('INFUSE Academy   ', 19, GREY, False, .04), run('01 / 03', 19, NAVY, True, .04)]])

# ---------- 02 ----------
s = new_slide('png/02-clean.png', TOOLS)
dash(s, 110, 110, MINT)
text(s, 166, 100, 600, 30, [[run('02  THE DARK FUNNEL', 22, MINT, True, .22)]])
text(s, 106, 560, 900, 130, [
    [run('of the buyer’s journey happens', 54, WHITE, True, -.015)],
    [run('where ', 54, WHITE, True, -.015), run('you cannot see.', 54, MINT, True, -.015)],
], lh=1.1)
text(s, 1100, 116, 500, 106, [[run('39%', 104, MINT, True, -.04)]], lh=1.0)
text(s, 1100, 228, 700, 34, [[run('is all your dashboard ever shows', 26, 'DCE3EA')]])
text(s, 1100, 276, 700, 26, [[run('PRICING PAGE · DEMO REQUEST · FORM FILL', 18, '5BC3AA', True, .16)]])
text(s, 889, 764, 320, 20, [[run('YOU START SEEING THEM HERE', 15, '62C9AE', True, .2)]], align=PP_ALIGN.CENTER)
text(s, 1560, 592, 320, 60, [[run('They finally', 22, 'DCE3EA')], [run('raise their hand.', 22, 'DCE3EA')]],
     lh=1.3, align=PP_ALIGN.CENTER)
text(s, 110, 1000, 500, 24, [[run('INFUSE Academy   ', 19, '7A8591', False, .04), run('02 / 03', 19, WHITE, True, .04)]])

# ---------- 03 ----------
s = new_slide('png/03-clean.png', TOOLS)
dash(s, 80, 68, ORANGE)
text(s, 136, 58, 500, 30, [[run('03  THE SCRIPT', 22, ORANGE, True, .22)]])
text(s, 1440, 60, 400, 24, [[run('INFUSE Academy   ', 19, GREY, False, .04), run('03 / 03', 19, NAVY, True, .04)]],
     align=PP_ALIGN.RIGHT)
text(s, 77, 96, 1800, 70, [[run('A small gap at the start. ', 56, NAVY, True, -.025),
                            run('A missed opportunity at scale.', 56, ORANGE, True, -.025)]], lh=1.1)
panels = [
    (80, '100%', NAVY, '5B6B7D', NAVY, 'The tool is capable. The logic is sound. The goal is compelling.'),
    (680, '38%', ORANGE, '5B6B7D', NAVY, 'Old habits creep back. Nobody raises a flag.'),
    (1280, '0%', 'FAFF81', '9AA3AF', WHITE, 'Nobody is using it. The gap was there on day one.'),
]
for x, big, big_c, sub_c, ink, cap in panels:
    text(s, x + 36, 580, 500, 106, [[run(big, 104, big_c, True, -.04), run('  team adoption', 22, sub_c)]], lh=1.0)
    text(s, x + 36, 698, 488, 70, [[run(cap, 25, ink)]], lh=1.3, wrap=True)

# ---------- 04 · process ----------
s = new_slide('png/04-clean.png', 'Process slide: tools used per visual and why AI.')
dash(s, 110, 94, ORANGE)
text(s, 166, 84, 500, 30, [[run('04  PROCESS', 22, ORANGE, True, .22)]])
text(s, 106, 124, 1000, 90, [[run('Tools & thinking', 72, NAVY, True, -.028)]], lh=1.1)
cards = [
    (110, 'A beginning', 'Magnific (character) + Claude → code',
     'The character was generated with AI; the world around her is code. The knot is built from hundreds of random loops so “hard” looks genuinely overwhelming. One orange thread escapes the chaos and ends right at her fingertips: the beginning is always within reach.'),
    (690, 'The dark funnel', 'Claude (Anthropic) → HTML Canvas',
     'A statistic is easy to read and easy to forget, so the 61% becomes physical: buyers flow through a funnel and most of their journey happens in the dark. In the live version your cursor is a flashlight, so you literally have to go looking for them.'),
    (1270, 'The script', 'Claude (Anthropic) → HTML + SVG',
     'The story is about a gap too small to notice. Each panel shows the team drifting away from the platform while one line underneath shows why: a few degrees off on day one becomes a missed opportunity by month six. The panels darken as the story does.'),
]
for x, title, tool, body in cards:
    text(s, x + 104, 306, 420, 44, [[run(title, 32, NAVY, True, -.02)]])
    text(s, x + 40, 384, 300, 22, [[run('TOOL', 16, ORANGE, True, .18)]])
    text(s, x + 40, 410, 470, 30, [[run(tool, 22, NAVY, True)]])
    text(s, x + 40, 462, 460, 240, [[run(body, 21, '3A4A5C')]], lh=1.45, wrap=True)
text(s, 158, 778, 300, 22, [[run('WHY AI', 16, MINT, True, .18)]])
text(s, 158, 810, 1604, 110, [[run('I generated the character in Magnific, using the INFUSE character sheet as the reference, because a human figure is where AI illustration shines and a flat, limited-palette prompt keeps it on brand. Everything around her — knot, thread, funnel, storyboard — is built in code with Claude, so it stays crisp, animates, and follows the palette exactly.', 24, WHITE)]], lh=1.42, wrap=True)

prs.save('INFUSE-Academy-visuals.pptx')

# PDF of the finished slides (text baked in) for the PDF submission route.
imgs = [Image.open(f'png/0{i}-full.png').convert('RGB') for i in (1, 2, 3, 4)]
imgs[0].save('INFUSE-Academy-visuals.pdf', save_all=True, append_images=imgs[1:], resolution=144)
print('ok')
